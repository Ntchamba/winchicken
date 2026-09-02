"""Excel (.xlsx) import for protocol rows — parse-only.

For non-technical farm users: they download a template, fill it in, and upload it on the
protocol configuration screen (onboarding step *and* the "Modifier" edit modal). This module
only **parses and validates** the file into a list of plain row dicts + a per-row report;
it never writes `ProtocolTemplate` rows. The frontend (`HouseProtocolForm`) merges the parsed
rows into the form so the user reviews everything before the normal save (see
docs/excel-import.md).

Decisions (docs/excel-import.md):
  * `openpyxl` (no pandas dependency).
  * Period unit for De/À is always "Jour" (the frontend applies DAY); "Unité" is a *resource*
    unit, only used when "Consommation" names a resource that must be created — an existing
    resource keeps its own unit and the file's value is ignored for that row.
  * Columns are matched by header name (case-insensitive, trimmed), not position.
  * A row missing "Catégorie", "De" or "Action", or with "À" < "De", is skipped and reported —
    one bad row never fails the whole import. Fully-empty rows are ignored silently.
  * "Quantité/jour" with no "Consommation" → the quantity is dropped, the row still imports.
  * A malformed "Créneaux" segment is dropped and reported as a *warning* — the row still
    imports, just without that time slot.
"""
import re
from io import BytesIO

from openpyxl import Workbook, load_workbook

# Exact French headers the template ships with.
HEADERS = ['Catégorie', 'De', 'À', 'Action', 'Détails', 'Consommation', 'Unité', 'Quantité/jour', 'Créneaux']
_HEADER_KEYS = {
    'catégorie': 'category',
    'de': 'from_value',
    'à': 'to_value',
    'action': 'what',
    'détails': 'details',
    'consommation': 'consumption',
    'unité': 'unit', 'unite': 'unit',
    'quantité/jour': 'quantity_per_day', 'quantite/jour': 'quantity_per_day',
    'créneaux': 'time_slots', 'creneaux': 'time_slots',
}

EXAMPLE_ROWS = [
    # a resource + unit + daily quantity, AND two time windows in one cell
    ['Alimentation', 1, 15, 'Aliment démarrage', '3000 kcal, 22,5% de protéines',
     'Provende', 'kg', 40, '06:30-07:30;18:30-19:30'],
    ['Alimentation', 15, 30, 'Aliment croissance', '3150 kcal, 21,5% de protéines', 'Provende', 'kg', 55, ''],
    ['Vaccination', 1, 1, 'Maladie de Newcastle', 'Hitchner B1, goutte oculaire', '', '', '', ''],
]

_SLOT_RE = re.compile(r'^\s*(\d{1,2}:\d{2})\s*-\s*(\d{1,2}:\d{2})\s*$')


def build_template_workbook() -> bytes:
    """A ready-to-edit .xlsx: the header row + a few realistic example rows."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Protocole'
    ws.append(HEADERS)
    for row in EXAMPLE_ROWS:
        ws.append(row)
    for i, header in enumerate(HEADERS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(14, len(header) + 4)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _clean(value):
    if value is None:
        return ''
    return str(value).strip()


def _as_number(value):
    """`value` -> float or None. Accepts "12", "12.5", "12,5" (French decimal), 12, 12.0."""
    if value is None or value == '':
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip().replace(',', '.'))
    except (TypeError, ValueError):
        return None


def _norm_time(hhmm):
    """'6:30' -> '06:30'; returns None if not a real HH:MM in range."""
    h, m = hhmm.split(':')
    h, m = int(h), int(m)
    if not (0 <= h <= 23 and 0 <= m <= 59):
        return None
    return f'{h:02d}:{m:02d}'


def _parse_slots(text):
    """Split a "Créneaux" cell on ';' and parse each `HH:MM-HH:MM` segment.
    Returns (slots, bad_segments): slots = [{'startTime','endTime'}]; bad_segments = raw strings
    that were malformed (unparseable or end <= start) and were dropped."""
    slots, bad = [], []
    for seg in _clean(text).split(';'):
        seg = seg.strip()
        if not seg:
            continue
        match = _SLOT_RE.match(seg)
        start = _norm_time(match.group(1)) if match else None
        end = _norm_time(match.group(2)) if match else None
        if not match or start is None or end is None or end <= start:
            bad.append(seg)
            continue
        slots.append({'startTime': start, 'endTime': end})
    return slots, bad


class ImportError(Exception):
    """Raised for a file that cannot be used at all (unreadable, no recognisable headers)."""


def parse_protocol_rows(file_obj) -> dict:
    """Parse an uploaded .xlsx into
    `{'rows': [...], 'imported': int, 'skipped': [...], 'warnings': [...]}`.

    Each ok row: {category, fromValue, toValue|None, untilEnd, what, details,
                  consumption|None, unit, quantityPerDay|None, timeSlots:[{startTime,endTime}]}.
    `skipped` / `warnings`: [{line, reason}] — `line` is the 1-based spreadsheet row number
    (header is line 1). `warnings` rows were imported (a bad Créneaux segment was dropped).
    """
    try:
        wb = load_workbook(file_obj, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a grab-bag of exceptions on a bad file
        raise ImportError('Fichier illisible. Utilisez un fichier .xlsx généré depuis le modèle.') from exc

    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise ImportError('Le fichier est vide.')

    col_index = {}
    for idx, cell in enumerate(header_row):
        key = _HEADER_KEYS.get(_clean(cell).lower())
        if key and key not in col_index:
            col_index[key] = idx

    missing = [h for h, k in (('Catégorie', 'category'), ('De', 'from_value'), ('Action', 'what'))
               if k not in col_index]
    if missing:
        raise ImportError(
            'En-têtes de colonnes introuvables : ' + ', '.join(f'« {h} »' for h in missing)
            + '. Téléchargez le modèle et conservez la première ligne.'
        )

    def cell(row, key):
        i = col_index.get(key)
        return row[i] if i is not None and i < len(row) else None

    rows, skipped, warnings = [], [], []
    for line, raw in enumerate(rows_iter, start=2):
        if raw is None or all(_clean(c) == '' for c in raw):
            continue  # fully-empty row — ignore, don't report

        category = _clean(cell(raw, 'category'))
        what = _clean(cell(raw, 'what'))
        from_raw = cell(raw, 'from_value')
        to_raw = cell(raw, 'to_value')

        if not category:
            skipped.append({'line': line, 'reason': 'catégorie manquante'})
            continue
        if not what:
            skipped.append({'line': line, 'reason': 'action manquante'})
            continue

        from_value = _as_number(from_raw)
        if from_value is None:
            skipped.append({'line': line, 'reason': (
                'valeur « De » manquante' if _clean(from_raw) == '' else f'valeur « De » invalide : « {_clean(from_raw)} »'
            )})
            continue

        until_end = _clean(to_raw) == ''
        to_value = None if until_end else _as_number(to_raw)
        if not until_end and to_value is None:
            skipped.append({'line': line, 'reason': f'valeur « À » invalide : « {_clean(to_raw)} »'})
            continue
        if to_value is not None and to_value < from_value:
            skipped.append({'line': line, 'reason': (
                f'« À » ({_clean(to_raw)}) est inférieur à « De » ({_clean(from_raw)})'
            )})
            continue

        consumption = _clean(cell(raw, 'consumption')) or None
        qpd = _as_number(cell(raw, 'quantity_per_day'))
        slots, bad_segments = _parse_slots(cell(raw, 'time_slots'))
        for seg in bad_segments:
            warnings.append({'line': line, 'reason': (
                f"créneau mal formé « {seg} » ignoré, ligne importée sans cet horaire"
            )})
        rows.append({
            'category': category,
            'fromValue': from_value,
            'toValue': to_value,
            'untilEnd': until_end,
            'what': what,
            'details': _clean(cell(raw, 'details')),
            'consumption': consumption,
            'unit': _clean(cell(raw, 'unit')),
            'quantityPerDay': qpd if consumption else None,
            'timeSlots': slots,
        })

    wb.close()
    return {'rows': rows, 'imported': len(rows), 'skipped': skipped, 'warnings': warnings}
