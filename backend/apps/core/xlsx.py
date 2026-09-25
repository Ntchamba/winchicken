"""Shared low-level helpers for the Excel (.xlsx) import features (Stock, Employees).

Only the generic spreadsheet plumbing lives here — cell cleaning, number coercion, header
matching by name, and the download response. The domain rules (category/supplier/role
matching, validation) stay in each feature's own module and reuse the existing serializers.
See docs/excel-import.md.
"""
import math
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook, load_workbook

CONTENT_TYPE_XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

TRUE_WORDS = {'oui', 'yes', 'true', 'vrai', '1', 'x', 'o'}

# A real import (employees or stock parameters for one farm) is kilobytes. Anything far larger is
# either a mistake or an attempt to exhaust memory while openpyxl parses it (security review
# 2026-09-26). read_only=True already streams rows, but a cap refuses the obvious bomb up front.
MAX_IMPORT_BYTES = 5 * 1024 * 1024


class WorkbookError(Exception):
    """A file that cannot be used at all (unreadable, empty, missing required headers, too big)."""


def check_upload_size(upload):
    """Raise WorkbookError (-> 400 with a French message) if the upload exceeds MAX_IMPORT_BYTES.
    `upload.size` is known before the file is read, so an oversized file is refused without
    parsing it."""
    size = getattr(upload, 'size', None)
    if size is not None and size > MAX_IMPORT_BYTES:
        raise WorkbookError(
            f'Fichier trop volumineux ({size // (1024 * 1024)} Mo). La taille maximale est de '
            f'{MAX_IMPORT_BYTES // (1024 * 1024)} Mo.'
        )


def clean(value):
    return '' if value is None else str(value).strip()


def as_number(value):
    """-> float or None. Accepts 12, 12.0, "12", "12.5", "12,5" (French decimal), "12 000". A cell reading
    "nan", "inf" or an overflowing "1e400" is not a number (float() would say it is)."""
    if value is None or value == '':
        return None
    if isinstance(value, bool):
        return None
    try:
        # Spaces, no-break and narrow no-break spaces group thousands in French: "12 000" is how
        # an FCFA amount is typed, and it used to read as unreadable (and import as 0).
        text = str(value).strip()
        for sep in (' ', '\u00a0', '\u202f'):
            text = text.replace(sep, '')
        number = float(value) if isinstance(value, (int, float)) else float(text.replace(',', '.'))
    except (TypeError, ValueError, OverflowError):
        return None
    return number if math.isfinite(number) else None


def as_bool(value):
    return clean(value).lower() in TRUE_WORDS


def open_rows(file_obj):
    """(header_row_tuple, iterator_of_data_rows). Raises WorkbookError on a bad/empty file."""
    try:
        wb = load_workbook(file_obj, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a grab-bag on a bad file
        raise WorkbookError('Fichier illisible. Utilisez un fichier .xlsx généré depuis le modèle.') from exc
    rows = wb.active.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration:
        raise WorkbookError('Le fichier est vide.')
    return header, rows


def header_index(header_row, aliases):
    """{internal_key: column_index} matching header cells by name (lower-cased, trimmed).
    `aliases` maps a normalised header string to an internal key."""
    index = {}
    for i, cell in enumerate(header_row):
        key = aliases.get(clean(cell).lower())
        if key and key not in index:
            index[key] = i
    return index


def require_headers(index, required):
    """`required`: list of (French label, internal key). Raises WorkbookError naming any missing."""
    missing = [label for label, key in required if key not in index]
    if missing:
        raise WorkbookError(
            'En-têtes de colonnes introuvables : ' + ', '.join(f'« {m} »' for m in missing)
            + '. Téléchargez le modèle et conservez la première ligne.'
        )


def cell_getter(index):
    def get(row, key):
        i = index.get(key)
        return row[i] if i is not None and i < len(row) else None
    return get


def build_workbook(sheet_title, headers, example_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(headers)
    for row in example_rows:
        ws.append(row)
    for i, header in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(14, len(str(header)) + 4)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_download(content, filename):
    resp = HttpResponse(content, content_type=CONTENT_TYPE_XLSX)
    resp['Content-Disposition'] = f'attachment; filename="{filename}"'
    return resp
